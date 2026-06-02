"""Export router — CSV и Excel."""
from __future__ import annotations

import calendar
import csv
import io
from datetime import date
from typing import Optional
from uuid import UUID

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.models.finance import Account, Category, Transaction
from app.infrastructure.api.v1.dependencies.auth import CanViewDashboard
from app.infrastructure.database.session import get_db
from app.services.report_service import (
    generate_cash_flow_report,
    generate_pl_report,
)

router = APIRouter(prefix="/companies/{company_id}/export", tags=["Экспорт"])

# ── Helpers ───────────────────────────────────────────────────────────────────

TYPE_LABELS = {"income": "Доход", "expense": "Расход", "transfer": "Перевод"}
STATUS_LABELS = {"confirmed": "Подтверждена", "draft": "Черновик", "reconciled": "Сверена"}


def _style_header_row(ws, row: int, n_cols: int):
    """Серый жирный заголовок."""
    fill = PatternFill("solid", fgColor="D9D9D9")
    font = Font(bold=True)
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 20


def _autofit(ws, min_w: int = 15, max_w: int = 40):
    """Автоподбор ширины колонок."""
    for col_cells in ws.columns:
        length = max(len(str(c.value or "")) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = max(min_w, min(max_w, length + 2))


MONEY_FMT = '#,##0.00 ₽'


async def _fetch_transactions(db: AsyncSession, company_id: UUID,
                              start: date, end: date, limit: int) -> list:
    result = await db.execute(
        select(
            Transaction.payment_date,
            Transaction.transaction_type,
            Transaction.amount,
            Transaction.currency,
            Transaction.description,
            Transaction.counterparty,
            Transaction.status,
            Category.name.label("category_name"),
            Account.name.label("account_name"),
        )
        .outerjoin(Category, Transaction.category_id == Category.id)
        .outerjoin(Account, Transaction.account_id == Account.id)
        .where(
            and_(
                Transaction.company_id == company_id,
                Transaction.is_deleted.is_(False),
                Transaction.payment_date >= start,
                Transaction.payment_date <= end,
            )
        )
        .order_by(Transaction.payment_date.desc())
        .limit(limit)
    )
    return result.all()


# ── CSV ───────────────────────────────────────────────────────────────────────

@router.get("/transactions/csv")
async def export_transactions_csv(
    company_id: UUID,
    start_date: Optional[date] = Query(default=None),
    end_date:   Optional[date] = Query(default=None),
    limit: int  = Query(default=1000, le=5000),
    current_user=Depends(CanViewDashboard),
    db: AsyncSession = Depends(get_db),
) -> StreamingResponse:
    today = date.today()
    start = start_date or date(today.year, today.month, 1)
    end   = end_date   or today

    rows = await _fetch_transactions(db, company_id, start, end, limit)

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Дата", "Тип", "Сумма", "Валюта", "Категория", "Счёт", "Контрагент", "Описание", "Статус"])
    for r in rows:
        writer.writerow([
            r.payment_date.strftime("%d.%m.%Y"),
            TYPE_LABELS.get(str(r.transaction_type.value if hasattr(r.transaction_type, 'value') else r.transaction_type), str(r.transaction_type)),
            str(r.amount).replace(".", ","),
            r.currency.value if hasattr(r.currency, 'value') else str(r.currency),
            r.category_name or "",
            r.account_name or "",
            r.counterparty or "",
            r.description or "",
            STATUS_LABELS.get(str(r.status.value if hasattr(r.status, 'value') else r.status), str(r.status)),
        ])

    content = "﻿" + output.getvalue()
    fname = f"transactions_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}.csv"
    return StreamingResponse(
        iter([content.encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ── Excel ─────────────────────────────────────────────────────────────────────

@router.get("/reports/xlsx")
async def export_reports_xlsx(
    company_id: UUID,
    start_date: Optional[date] = Query(default=None),
    end_date:   Optional[date] = Query(default=None),
    current_user=Depends(CanViewDashboard),
    db: AsyncSession = Depends(get_db),
) -> StreamingResponse:
    today = date.today()
    start = start_date or date(today.year, 1, 1)
    end   = end_date   or today

    # ── Данные ──────────────────────────────────────────────────────────────
    pl       = await generate_pl_report(db, company_id, start, end)
    cashflow = await generate_cash_flow_report(db, company_id, start, end)
    rows_txn = await _fetch_transactions(db, company_id, start, end, 500)

    # Последние 6 месяцев
    monthly_rows = []
    for i in range(5, -1, -1):
        total = today.year * 12 + (today.month - 1) - i
        y, m = total // 12, (total % 12) + 1
        ms = date(y, m, 1)
        me = today if (y == today.year and m == today.month) else date(y, m, calendar.monthrange(y, m)[1])
        mp = await generate_pl_report(db, company_id, ms, me)
        monthly_rows.append({
            "label": ms.strftime("%b %Y"),
            "revenue": mp.get("revenue", 0),
            "expenses": mp.get("opex", 0),
            "net_income": mp.get("net_income", 0),
            "ebitda": mp.get("ebitda", 0),
        })

    # ── Книга ───────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()

    # === Лист 1: P&L ===
    ws_pl = wb.active
    ws_pl.title = "P&L"
    ws_pl.append(["Показатель", "Значение, ₽"])
    _style_header_row(ws_pl, 1, 2)
    for label, key in [
        ("Выручка", "revenue"), ("Себестоимость", "cogs"), ("Валовая прибыль", "gross_profit"),
        ("Операционные расходы", "opex"), ("EBITDA", "ebitda"), ("Налоги", "taxes"), ("Чистая прибыль", "net_income"),
    ]:
        row = ws_pl.append([label, pl.get(key, 0)])
        ws_pl.cell(ws_pl.max_row, 2).number_format = MONEY_FMT
    _autofit(ws_pl)

    # === Лист 2: Cash Flow ===
    ws_cf = wb.create_sheet("Cash Flow")
    ws_cf.append(["Показатель", "Значение, ₽"])
    _style_header_row(ws_cf, 1, 2)
    for label, key in [
        ("Операционный CF", "operating_cf"), ("Инвестиционный CF", "investing_cf"),
        ("Финансовый CF", "financing_cf"), ("Чистый денежный поток", "net_cash_flow"),
        ("Баланс на начало", "opening_balance"), ("Баланс на конец", "closing_balance"),
    ]:
        ws_cf.append([label, cashflow.get(key, 0)])
        ws_cf.cell(ws_cf.max_row, 2).number_format = MONEY_FMT
    _autofit(ws_cf)

    # === Лист 3: По месяцам ===
    ws_m = wb.create_sheet("По месяцам")
    ws_m.append(["Месяц", "Доходы", "Расходы", "Прибыль", "EBITDA"])
    _style_header_row(ws_m, 1, 5)
    for mr in monthly_rows:
        ws_m.append([mr["label"], mr["revenue"], mr["expenses"], mr["net_income"], mr["ebitda"]])
        for col in range(2, 6):
            ws_m.cell(ws_m.max_row, col).number_format = MONEY_FMT
    ws_m.freeze_panes = "A2"
    _autofit(ws_m)

    # === Лист 4: Транзакции ===
    ws_t = wb.create_sheet("Транзакции")
    ws_t.append(["Дата", "Тип", "Сумма", "Валюта", "Категория", "Счёт", "Контрагент", "Описание", "Статус"])
    _style_header_row(ws_t, 1, 9)
    for r in rows_txn:
        txn_type = str(r.transaction_type.value if hasattr(r.transaction_type, 'value') else r.transaction_type)
        status   = str(r.status.value if hasattr(r.status, 'value') else r.status)
        currency = r.currency.value if hasattr(r.currency, 'value') else str(r.currency)
        ws_t.append([
            r.payment_date.strftime("%d.%m.%Y"),
            TYPE_LABELS.get(txn_type, txn_type),
            float(r.amount),
            currency,
            r.category_name or "",
            r.account_name or "",
            r.counterparty or "",
            r.description or "",
            STATUS_LABELS.get(status, status),
        ])
        ws_t.cell(ws_t.max_row, 3).number_format = MONEY_FMT
    ws_t.freeze_panes = "A2"
    _autofit(ws_t)

    # ── Сохраняем и отдаём ──────────────────────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = f"planfact_report_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
